#!/usr/bin/env python3
"""Monta a planilha de acompanhamento: um painel, uma aba por materia e as aulas do Gran.

O arquivo sai em planilha/TCDF-acompanhamento.xlsx. Solto no Google Drive, ele
vira uma planilha Google com as abas e as formulas vivas.

Uso:
    python3 scripts/gerar_planilha.py
"""
import json, re, sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ / "scripts"))
import aulas as mapa_aulas
SAIDA = RAIZ / "planilha" / "TCDF-acompanhamento.xlsx"

ROXO = "6D28D9"; LILAS = "F5F3FF"; CINZA = "6B7280"; ZEBRA = "FAFAFA"
AMARELO = "FFF9DB"; BORDA = "D1D5DB"
FONTE = "Arial"

ABAS = {
    "administracao-financeira-orcamentaria": "AFO",
    "nocoes-direito-tributario": "D. Tributário",
    "direito-administrativo": "D. Administrativo",
    "gestao-de-contratos": "Gestão de Contratos",
    "direito-constitucional": "D. Constitucional",
    "lei-organica-regimento-tcdf": "LO e RI do TCDF",
    "lei-organica-df": "LODF",
    "administracao-geral-e-publica": "Adm. Geral e Pública",
    "analise-dados-estatistica-ia": "Análise de Dados",
    "regime-juridico-servidores-df": "LC 840 Servidores",
    "direito-previdenciario": "D. Previdenciário",
    "nocoes-direito-civil": "D. Civil",
    "lingua-portuguesa": "Português",
    "raciocinio-logico-matematica-financeira": "RLM",
    "conhecimentos-df-politica-mulheres": "DF e Mulheres",
    "nocoes-primeiros-socorros": "Primeiros Socorros",
}
N_FILTROS = 8
N_REGISTROS = 20
FINA = Side(style="thin", color=BORDA)
GRADE = Border(left=FINA, right=FINA, top=FINA, bottom=FINA)


def partes(texto):
    m = re.match(r"^(\d+)\s+(.*)$", texto.strip(), re.S)
    numero, resto = (m.group(1), m.group(2)) if m else ("", texto)
    if ":" in resto:
        titulo, cauda = resto.split(":", 1)
        subs = [s.strip() for s in re.split(r";\s*(?=\d+\.\d)", cauda.strip()) if s.strip()]
        if len(subs) == 1 and not re.match(r"^\d+\.\d", subs[0]):
            return numero, resto.strip(), []
        return numero, titulo.strip(), subs
    return numero, resto.strip(), []


def cabecalho(ws, celula, texto, largura=None):
    c = ws[celula]
    c.value = texto
    c.font = Font(name=FONTE, bold=True, color="FFFFFF", size=9)
    c.fill = PatternFill("solid", fgColor=ROXO)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def main():
    ed = json.loads((RAIZ / "scripts/edital.json").read_text(encoding="utf-8"))
    prog = json.loads((RAIZ / "scripts/progressao.json").read_text(encoding="utf-8"))
    gran = json.loads((RAIZ / "scripts/gran.json").read_text(encoding="utf-8"))
    mats = {m["id"]: m for m in ed["materias"]}

    mapa = mapa_aulas.mapa_de_aulas(gran, mats)

    def aula_de(mid, chave):
        """Numero da aula; entre parenteses quando vem do topico, nao do subtopico."""
        m = mapa.get((mid, chave))
        if not m:
            return ""
        lista = ", ".join(str(x) for x in m["aulas"])
        return lista if m["precisao"] == mapa_aulas.EXATA else f"({lista})"

    wb = Workbook()
    painel = wb.active
    painel.title = "Painel"

    ordem, dia_da = [], {}
    for g in prog["grupos"]:
        for b in g["materias"]:
            ordem.append(b["id"])
            dia_da[b["id"]] = g["dia"]

    posicoes = {}
    for mid in ordem:
        info = mats[mid]
        ws = wb.create_sheet(ABAS[mid])
        ws.sheet_properties.tabColor = ROXO

        ws["A1"] = info["nome"]
        ws["A1"].font = Font(name=FONTE, bold=True, size=14, color=ROXO)
        ws.merge_cells("A1:D1")
        ws["E1"] = (f"Bloco {info['bloco']} · {dia_da[mid]} · {len(info['ementa'])} tópicos · "
                    f"{info['itens_estimados']} itens estimados do edital")
        ws["E1"].font = Font(name=FONTE, size=9, color=CINZA)

        ws["A2"] = ("Escreva o nome de cada filtro na linha 4 e marque X nas linhas do que ele sorteia. "
                    "Amarelo = você preenche. Aula entre parênteses = a aula cobre o tópico, "
                    "mas o rótulo dela não cita esse subtópico.")
        ws["A2"].font = Font(name=FONTE, size=9, italic=True, color=CINZA)
        ws.merge_cells("A2:M2")

        ws["D4"] = "Nome do filtro no Gran →"
        ws["D4"].font = Font(name=FONTE, bold=True, size=9)
        ws["D4"].alignment = Alignment(horizontal="right")
        ws["D5"] = "Este filtro venceu (≥90%)? →"
        ws["D5"].font = Font(name=FONTE, bold=True, size=9)
        ws["D5"].alignment = Alignment(horizontal="right")
        for i in range(N_FILTROS):
            col = get_column_letter(5 + i)
            for linha in (4, 5):
                c = ws[f"{col}{linha}"]
                c.fill = PatternFill("solid", fgColor=AMARELO)
                c.alignment = Alignment(horizontal="center")
        ws["M4"] = "GERAL N2"
        ws["M4"].font = Font(name=FONTE, bold=True, size=8, color=ROXO)
        ws["M4"].alignment = Alignment(horizontal="center")
        ws["M5"].fill = PatternFill("solid", fgColor=AMARELO)
        ws["M5"].alignment = Alignment(horizontal="center")

        L_CAB = 7
        cabecalho(ws, f"A{L_CAB}", "Nº")
        cabecalho(ws, f"B{L_CAB}", "Tópico e subtópico do edital")
        cabecalho(ws, f"C{L_CAB}", "Peso")
        cabecalho(ws, f"D{L_CAB}", "Aula Gran")
        for i in range(N_FILTROS):
            cabecalho(ws, f"{get_column_letter(5 + i)}{L_CAB}", f"F{i + 1}")
        cabecalho(ws, f"M{L_CAB}", "GERAL N2")

        linha = L_CAB + 1
        linhas_topico = []
        for t in info["ementa"]:
            n = t["n"]
            numero, titulo, subs = partes(t["texto"])
            ws[f"A{linha}"] = numero or n
            ws[f"B{linha}"] = titulo
            ws[f"C{linha}"] = t["peso"]
            ws[f"D{linha}"] = aula_de(mid, n)
            for col in "ABCD":
                ws[f"{col}{linha}"].font = Font(name=FONTE, bold=True, size=10)
            linhas_topico.append(linha)
            for col in "ABCDEFGHIJKL":
                ws[f"{col}{linha}"].fill = PatternFill("solid", fgColor=LILAS)
            ws[f"M{linha}"].fill = PatternFill("solid", fgColor=AMARELO)
            linha += 1
            itens = subs + [f"{s} (subdivisão sugerida)" for s in t.get("subtopicos_sugeridos", [])]
            for s in itens:
                ws[f"B{linha}"] = "    " + s
                ws[f"B{linha}"].font = Font(name=FONTE, size=9, color="374151")
                chave = s.split()[0]
                ws[f"D{linha}"] = aula_de(mid, chave)
                ws[f"D{linha}"].font = Font(
                    name=FONTE, size=9,
                    color=CINZA if str(ws[f"D{linha}"].value).startswith("(") else ROXO)
                linha += 1
        centro = Alignment(horizontal="center")
        for lin in range(L_CAB + 1, linha):
            for col in "ACD":
                ws[f"{col}{lin}"].alignment = centro
            ws[f"B{lin}"].alignment = Alignment(wrap_text=True, vertical="top")

        # ---- registro dos simulados ----
        reg = linha + 2
        ws[f"A{reg}"] = "REGISTRO DOS SIMULADOS DESTA MATÉRIA"
        ws[f"A{reg}"].font = Font(name=FONTE, bold=True, size=11, color=ROXO)
        cab = ["Filtro", "Data", "Total", "Acertos", "Erros", "% Bruto", "% Líquido", "Veredito"]
        for i, texto in enumerate(cab):
            cabecalho(ws, f"{get_column_letter(1 + i)}{reg + 1}", texto)
        # Linha de exemplo fica FORA do bloco de dados: ensina o formato sem
        # entrar nas contas do painel.
        ex = reg + 2
        ws[f"A{ex}"] = "exemplo — apague ou ignore"
        ws[f"B{ex}"] = "15/09/2026"
        ws[f"C{ex}"] = 20
        ws[f"D{ex}"] = 18
        ws[f"E{ex}"] = 2
        ws[f"F{ex}"] = f'=IFERROR(D{ex}/C{ex},"")'
        ws[f"G{ex}"] = f'=IFERROR((D{ex}-E{ex})/C{ex},"")'
        ws[f"H{ex}"] = f'=IF(N(C{ex})=0,"",IF(D{ex}/C{ex}>=0.9,"VENCEU","REPETE"))'
        for col in "ABCDEFGH":
            ws[f"{col}{ex}"].font = Font(name=FONTE, size=9, italic=True, color=CINZA)
        ws[f"F{ex}"].number_format = "0.0%"
        ws[f"G{ex}"].number_format = "0.0%"

        primeira = ex + 1
        for k in range(N_REGISTROS):
            lin = primeira + k
            for i in range(5):
                ws[f"{get_column_letter(1 + i)}{lin}"].fill = PatternFill("solid", fgColor=AMARELO)
            ws[f"F{lin}"] = f'=IFERROR(D{lin}/C{lin},"")'
            ws[f"G{lin}"] = f'=IFERROR((D{lin}-E{lin})/C{lin},"")'
            ws[f"H{lin}"] = f'=IF(N(C{lin})=0,"",IF(D{lin}/C{lin}>=0.9,"VENCEU","REPETE"))'
            ws[f"F{lin}"].number_format = "0.0%"
            ws[f"G{lin}"].number_format = "0.0%"
            ws[f"H{lin}"].alignment = centro
        posicoes[mid] = {"aba": ABAS[mid], "reg_ini": primeira, "reg_fim": primeira + N_REGISTROS - 1,
                         "n_topicos": len(info["ementa"])}

        ws.column_dimensions["A"].width = 5
        ws.column_dimensions["B"].width = 62
        ws.column_dimensions["C"].width = 6
        ws.column_dimensions["D"].width = 10
        for i in range(N_FILTROS):
            ws.column_dimensions[get_column_letter(5 + i)].width = 5
        ws.column_dimensions["M"].width = 10
        ws.freeze_panes = f"A{L_CAB + 1}"

    # ---------- painel ----------
    painel["A1"] = "TCDF / ANACE 2026 — acompanhamento por matéria"
    painel["A1"].font = Font(name=FONTE, bold=True, size=16, color=ROXO)
    painel["A2"] = ("Uma aba por matéria. Nesta folha, tudo é fórmula: os números se atualizam sozinhos "
                    "conforme você preenche as abas. Meta: 90% bruto.")
    painel["A2"].font = Font(name=FONTE, size=10, italic=True, color=CINZA)
    painel.merge_cells("A2:H2")

    cab = ["Dia", "Matéria", "Tópicos", "Itens do edital", "Filtros vencidos",
           "Lotes registrados", "Melhor % bruto", "Último % bruto"]
    for i, texto in enumerate(cab):
        cabecalho(painel, f"{get_column_letter(1 + i)}4", texto)

    lin = 5
    for mid in ordem:
        p = posicoes[mid]
        aba = f"'{p['aba']}'"
        painel[f"A{lin}"] = dia_da[mid]
        painel[f"B{lin}"] = mats[mid]["nome"]
        painel[f"C{lin}"] = p["n_topicos"]
        painel[f"D{lin}"] = mats[mid]["itens_estimados"]
        painel[f"E{lin}"] = f'=COUNTIF({aba}!E5:L5,"S")'
        painel[f"F{lin}"] = f'=COUNT({aba}!C{p["reg_ini"]}:C{p["reg_fim"]})'
        painel[f"G{lin}"] = f'=IFERROR(MAX({aba}!F{p["reg_ini"]}:F{p["reg_fim"]}),"")'
        painel[f"H{lin}"] = (f'=IFERROR(INDEX({aba}!F{p["reg_ini"]}:F{p["reg_fim"]},'
                             f'COUNT({aba}!C{p["reg_ini"]}:C{p["reg_fim"]})),"")')
        for i in range(8):
            c = painel[f"{get_column_letter(1 + i)}{lin}"]
            c.font = Font(name=FONTE, size=10)
            if i != 1:
                c.alignment = Alignment(horizontal="center")
        painel[f"G{lin}"].number_format = "0.0%"
        painel[f"H{lin}"].number_format = "0.0%"
        lin += 1

    total = lin
    painel[f"B{total}"] = "TOTAL"
    painel[f"C{total}"] = f"=SUM(C5:C{total - 1})"
    painel[f"D{total}"] = f"=SUM(D5:D{total - 1})"
    painel[f"E{total}"] = f"=SUM(E5:E{total - 1})"
    painel[f"F{total}"] = f"=SUM(F5:F{total - 1})"
    for i in range(8):
        c = painel[f"{get_column_letter(1 + i)}{total}"]
        c.font = Font(name=FONTE, bold=True, size=10)
        c.fill = PatternFill("solid", fgColor=LILAS)
        if i != 1:
            c.alignment = Alignment(horizontal="center")

    nota = total + 2
    painel[f"A{nota}"] = "COMO USAR"
    painel[f"A{nota}"].font = Font(name=FONTE, bold=True, size=11, color=ROXO)
    for i, texto in enumerate([
        "1. Na aba da matéria, escreva na linha 4 o nome de cada filtro que você montar no Gran (F1 a F8).",
        "2. Marque X na coluna do filtro, na linha de cada tópico e subtópico que ele sorteia.",
        "3. Resolveu um lote? Lance filtro, data, total, acertos e erros no bloco de baixo. "
        "As colunas % Bruto, % Líquido e Veredito são fórmulas — não digite nada nelas.",
        "4. Filtro que bateu 90%: escreva S na linha 5 e jogue o filtro dentro do GERAL N2, "
        "o filtro único da matéria vencida, de onde saem os simulados de grupo e as revisões.",
        "5. Tópico vencido (18/20 sobre a ementa inteira dele): marque X na coluna GERAL N2.",
        "",
        "Células amarelas são suas. O resto é fórmula ou vem do edital.",
        "Bruto = acertos ÷ total, é ele que decide o avanço. Líquido = (acertos − erros) ÷ total, "
        "é a régua da prova, onde cada erro anula um acerto. Branco conta como erro.",
    ]):
        painel[f"A{nota + 1 + i}"] = texto
        painel[f"A{nota + 1 + i}"].font = Font(name=FONTE, size=9,
                                               color=CINZA if i >= 6 else "1A1A1A")
        painel.merge_cells(f"A{nota + 1 + i}:H{nota + 1 + i}")

    painel.column_dimensions["A"].width = 16
    painel.column_dimensions["B"].width = 44
    for col in "CDEFGH":
        painel.column_dimensions[col].width = 15
    painel.freeze_panes = "A5"

    # ---------- aulas do Gran ----------
    wa = wb.create_sheet("Aulas do Gran")
    wa["A1"] = "Aulas do curso do Gran e o que cada uma cobre do edital"
    wa["A1"].font = Font(name=FONTE, bold=True, size=14, color=ROXO)
    wa["A2"] = ("Extraído do rótulo de cada aula. O curso fatia e reordena o programa, então a "
                "numeração das aulas não segue a do edital.")
    wa["A2"].font = Font(name=FONTE, size=10, italic=True, color=CINZA)
    wa.merge_cells("A2:D2")
    for i, texto in enumerate(["Matéria", "Aula", "O que a aula cobre", "Tópicos do edital"]):
        cabecalho(wa, f"{get_column_letter(1 + i)}4", texto)
    lin = 5
    for mid in ordem:
        for a in sorted(gran["materias"].get(mid, []), key=lambda x: x["aula"]):
            wa[f"A{lin}"] = mats[mid]["nome"]
            wa[f"B{lin}"] = a["aula"]
            wa[f"C{lin}"] = a.get("rotulo", "")
            wa[f"D{lin}"] = ", ".join(a["topicos"]) or "—"
            for col in "ABCD":
                wa[f"{col}{lin}"].font = Font(name=FONTE, size=9)
            wa[f"B{lin}"].alignment = Alignment(horizontal="center")
            lin += 1
    wa.column_dimensions["A"].width = 38
    wa.column_dimensions["B"].width = 7
    wa.column_dimensions["C"].width = 70
    wa.column_dimensions["D"].width = 16
    wa.freeze_panes = "A5"

    SAIDA.parent.mkdir(exist_ok=True)
    wb.save(SAIDA)
    print(f"{SAIDA} · {len(wb.sheetnames)} abas: {', '.join(wb.sheetnames[:5])}...")


if __name__ == "__main__":
    main()
